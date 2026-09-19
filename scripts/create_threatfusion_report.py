"""Build the plain-language ThreatFusion report."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "reports" / "ThreatFusion_Simple_Workflow_Report.docx"
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
MUTED = RGBColor(89, 104, 120)


def set_font(run, size=11, color=None, bold=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]):
    tbl_pr = table._tbl.tblPr
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_layout = OxmlElement("w:tblLayout")
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_pr.append(tbl_layout)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)


def body_paragraph(doc, text: str, before=0, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    set_font(p.add_run(text))
    return p


def heading(doc, text: str, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    run = p.add_run(text)
    set_font(run, 16 if level == 1 else 13, BLUE if level == 1 else DARK, True)
    return p


def item(doc, label: str, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.1
    set_font(p.add_run("• "), color=BLUE, bold=True)
    set_font(p.add_run(label + ": "), bold=True)
    set_font(p.add_run(text))


def main():
    OUT.parent.mkdir(exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("ThreatFusion | Simple Workflow Report"), 9, MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("Internal demo report | Synthetic telemetry only"), 9, MUTED)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("ThreatFusion: How It Works"), 24, DARK, True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_font(subtitle.add_run("A simple explanation of defense telemetry, actionable incidents, and response guidance"), 13, MUTED)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(14)
    set_font(meta.add_run("Purpose: "), 10, bold=True)
    set_font(meta.add_run("Explain how the system turns many security alerts into a small number of incidents the defense team can act on."), 10)

    heading(doc, "What the system does")
    body_paragraph(doc, "ThreatFusion is a defense operations tool. It takes alerts from security tools, looks for events that belong together, and presents only the cases that have enough evidence to justify action.")
    body_paragraph(doc, "It does not automatically block systems or claim to identify an attacker. A human analyst reviews the evidence and approves any response.")

    heading(doc, "The simple workflow")
    steps = [
        ("1. Ingest", "Receive telemetry from CTI reports, SIEM logs, endpoint tools, and network sensors."),
        ("2. Connect", "Group alerts that share a host, user, IP address, or confirmed indicator and occur close together in time."),
        ("3. Understand", "Map suspicious behavior to MITRE ATT&CK techniques and check whether the activity follows a believable attack sequence."),
        ("4. Validate", "Score the evidence, check the number and independence of sources, and reduce confidence when benign context is present."),
        ("5. Respond", "Promote only validated cases to actionable incidents, then show the analyst a short runbook and evidence summary."),
    ]
    for label, text in steps:
        item(doc, label, text)

    heading(doc, "When an incident becomes actionable")
    body_paragraph(doc, "A candidate is promoted only when every check below is met. This prevents a single noisy alert from becoming an incident.")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    headers = ["Check", "Required benchmark"]
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, "E8EEF5")
        r = cell.paragraphs[0].add_run(text)
        set_font(r, 10, DARK, True)
    rows = [
        ("Behavior evidence", "At least 2 mapped ATT&CK technique events"),
        ("Attack progression", "At least 2 tactics and an attack-flow score of 0.50 or above"),
        ("Evidence confidence", "55% or above after corroboration and any contradiction penalties"),
        ("Source independence", "50% or above, so one repeated source does not dominate the result"),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, (label, value)):
            r = cell.paragraphs[0].add_run(text)
            set_font(r, 10)
    set_table_geometry(table, [2600, 6760])

    heading(doc, "How priority is decided")
    body_paragraph(doc, "After promotion, the system assigns a priority score. The score is 45% evidence confidence, 25% threat severity, 20% mission impact, and 10% urgency. A P1 is 82 or higher; P2 is 65 or higher.")
    item(doc, "Confidence", "Do several independent sources and techniques support the same case?")
    item(doc, "Severity", "How risky are the observed techniques?")
    item(doc, "Mission impact", "Does the case involve an important business or mission system?")
    item(doc, "Urgency", "Does the activity look active and fast-moving?")

    heading(doc, "What the analyst sees")
    body_paragraph(doc, "For each actionable incident, the console provides a plain-language summary, observed timeline, supporting and weakening evidence, risk score, affected assets, and recommended defensive actions.")
    item(doc, "Defense feed", "The bundled demo can ingest 18 production-shaped synthetic records. The sample includes normal activity as well as one correlated intrusion sequence.")
    item(doc, "IBM Bob", "Bob answers questions about the selected case using only grounded case facts. It gives concise answers about priority, evidence, gaps, and the incident runbook.")

    heading(doc, "Important limits")
    body_paragraph(doc, "This prototype is designed for demonstration and analyst support. Its bundled data is synthetic, and its offline benchmark is not a production accuracy claim. A real deployment needs customer-approved telemetry, larger labeled datasets, ongoing threshold calibration, and analyst review.")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
